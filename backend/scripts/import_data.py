import argparse
import re
import sys
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete, inspect, text

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))

from app.database import Base, SessionLocal, engine  # noqa: E402
from app.models import Food, FoodNutrient, ProductStandardMapping, StandardFood  # noqa: E402
from app.services.mfds import MfdsClient  # noqa: E402
from app.services.standard_foods import estimate_serving, match_standard_food  # noqa: E402

NUTRIENT_COLUMNS = {
    "kcal": (6, "kcal"), "protein": (8, "g"), "calcium": (22, "mg"),
    "iron": (23, "mg"), "vitaminA": (34, "µgRAE"), "vitaminC": (51, "mg"),
    "sodium": (27, "mg"),
}

MFDS_PRODUCTS = [
    {"cu_id": "5228", "mfds_query": "바나나킥", "mfds_maker": "농심"},
]

MFDS_BY_ID = {item["cu_id"]: item for item in MFDS_PRODUCTS}


def existing_file(value: str) -> Path:
    path = Path(value).expanduser().resolve()
    if not path.is_file():
        raise argparse.ArgumentTypeError(f"파일을 찾을 수 없습니다: {path}")
    return path


def number(value: object) -> float | None:
    if value is None or str(value).strip() in {"", "-", "N/A", "Tr"}:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def extract_barcode(url: str | None) -> str | None:
    match = re.search(r"(\d{13})(?:_\d+)?\.(?:jpg|png)$", url or "", re.IGNORECASE)
    return match.group(1) if match else None


def split_candidates(value: object) -> list[str]:
    return [item.strip() for item in str(value or "").split(";") if item.strip()]


def resolve_item_report(source: dict, mfds) -> tuple[str | None, str | None, str | None]:
    if mfds:
        return (
            mfds.item_report_no,
            "식약처 API 확인",
            f"식약처 영양성분 API 제품명과 보고번호 확인: {mfds.item_report_no}",
        )

    candidates = source["item_report_candidates"]
    if not candidates:
        return None, source["item_report_status"], source["item_report_evidence"]

    evidence = str(source["item_report_evidence"] or "제품명 유사도 기준 후보")
    if len(candidates) > 1:
        evidence = f"{evidence}. 복수 후보 중 첫 번째 번호를 대표값으로 자동 확정"
    else:
        evidence = f"{evidence}. 단일 후보를 자동 확정"
    return candidates[0], "상품명 유사 자동 확정", evidence


def load_cu_rows(path: Path) -> dict[str, dict]:
    ws = load_workbook(path, read_only=True, data_only=True)["CU 식품·음료"]
    headers = [str(cell.value or "").strip() for cell in ws[7]]
    columns = {header: index for index, header in enumerate(headers)}
    required = ["분류", "상품명", "가격(원)", "CU 상품 ID", "이미지 URL"]
    missing = [header for header in required if header not in columns]
    if missing:
        raise RuntimeError(f"CU 상품 파일에 필요한 열이 없습니다: {', '.join(missing)}")

    def cell(row: tuple, header: str) -> object:
        index = columns.get(header)
        return row[index] if index is not None and index < len(row) else None

    result = {}
    for row in ws.iter_rows(min_row=8, values_only=True):
        cu_id = str(cell(row, "CU 상품 ID") or "").strip()
        if not cu_id:
            continue
        image_url = cell(row, "이미지 URL")
        candidate_value = cell(row, "품목제조보고번호 (후보)")
        result[cu_id] = {
            "category": cell(row, "분류"),
            "name": str(cell(row, "상품명") or "").strip(),
            "price": int(cell(row, "가격(원)")) if cell(row, "가격(원)") is not None else None,
            "image_url": image_url,
            "barcode": extract_barcode(image_url),
            "item_report_candidates": split_candidates(candidate_value),
            "item_report_status": cell(row, "대조 상태"),
            "item_report_evidence": cell(row, "대조 근거"),
        }
    return result


def load_standard_rows(path: Path) -> dict[str, dict]:
    ws = load_workbook(path, read_only=True, data_only=True)["국가표준식품성분 Database 10.4"]
    result = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        code = str(row[0] or "")
        if code:
            result[code] = {
                "name": str(row[3]), "food_group": str(row[2]),
                "nutrients": {key: number(row[column - 1]) for key, (column, _) in NUTRIENT_COLUMNS.items()},
            }
    return result


def mfds_match(config: dict):
    if not config.get("mfds_query"):
        return None
    try:
        matches = MfdsClient().search(config["mfds_query"], config.get("mfds_maker"))
    except Exception as error:
        print(f"식약처 조회를 건너뜁니다: {type(error).__name__}")
        return None
    return next((item for item in matches if item.item_report_no == "199104611019"), None) or (matches[0] if matches else None)


def ensure_food_candidate_columns() -> None:
    """create_all이 기존 개발 DB에 추가하지 못하는 열을 가볍게 보완한다."""
    existing = {column["name"] for column in inspect(engine).get_columns("foods")}
    definitions = {
        "item_report_candidates": "JSON",
        "item_report_status": "VARCHAR(80)",
        "item_report_evidence": "TEXT",
    }
    with engine.begin() as connection:
        for name, sql_type in definitions.items():
            if name not in existing:
                connection.execute(text(f"ALTER TABLE foods ADD COLUMN {name} {sql_type}"))


def seed(cu_path: Path, standard_path: Path) -> None:
    cu_rows = load_cu_rows(cu_path)
    standard_rows = load_standard_rows(standard_path)
    missing_mfds = set(MFDS_BY_ID) - set(cu_rows)
    if missing_mfds:
        raise RuntimeError(f"식약처 연결 대상 CU 상품을 찾지 못했습니다: {', '.join(sorted(missing_mfds))}")

    Base.metadata.create_all(bind=engine)
    ensure_food_candidate_columns()
    with SessionLocal() as db:
        db.execute(delete(ProductStandardMapping).where(ProductStandardMapping.food_id.like("cu:%")))
        db.execute(delete(FoodNutrient).where(FoodNutrient.food_id.like("cu:%")))
        db.execute(delete(Food).where(Food.id.like("cu:%")))

        used_barcodes: set[str] = set()
        enriched_names: list[str] = []
        added_standard_codes: set[str] = set()
        parsed_serving_count = 0
        for cu_id, source in cu_rows.items():
            config = MFDS_BY_ID.get(cu_id)
            mfds = mfds_match(config) if config else None
            if mfds and source["item_report_candidates"] and mfds.item_report_no not in source["item_report_candidates"]:
                print(f"식약처 번호가 엑셀 후보와 달라 영양 연결을 건너뜁니다: {source['name']}")
                mfds = None
            barcode = source["barcode"]
            if barcode in used_barcodes:
                barcode = None
            elif barcode:
                used_barcodes.add(barcode)
            item_report_no, item_report_status, item_report_evidence = resolve_item_report(
                source, mfds
            )
            serving = estimate_serving(source["name"], source["category"])
            standard_rule = (
                match_standard_food(source["name"], source["category"])
                if item_report_no else None
            )
            standard_code = standard_rule.food_code if standard_rule else None
            standard = standard_rows.get(standard_code)
            if standard_rule and standard is None:
                raise RuntimeError(f"국가표준식품 코드가 없습니다: {standard_code}")
            food = Food(
                id=f"cu:{cu_id}", source_type="CU_PRODUCT", name=source["name"],
                brand=source["name"].split(")", 1)[0] if ")" in source["name"] else "CU",
                category=source["category"], price=source["price"], barcode=barcode,
                image_url=source["image_url"],
                item_report_no=item_report_no,
                item_report_candidates=source["item_report_candidates"],
                item_report_status=item_report_status,
                item_report_evidence=item_report_evidence,
                serving_amount=Decimal(str(serving.amount)) if standard or mfds else None,
                serving_unit=serving.unit if standard or mfds else None,
                count_unit="회" if standard or mfds else None,
                serving_label=serving.label if standard or mfds else None,
            )
            db.add(food)

            if standard:
                if standard_code not in added_standard_codes and db.get(StandardFood, standard_code) is None:
                    db.add(StandardFood(food_code=standard_code, name=standard["name"],
                                        food_group=standard["food_group"], nutrients_per_100g=standard["nutrients"]))
                added_standard_codes.add(standard_code)
                db.add(ProductStandardMapping(food_id=food.id, standard_food_code=standard_code,
                                               match_method=("CATEGORY_DEFAULT" if standard_rule.confidence <= 0.55
                                                             else "NORMALIZED_RULE"),
                                               confidence=Decimal(str(standard_rule.confidence)),
                                               review_status="AUTO_MATCHED"))

            if not mfds and not standard:
                continue
            enriched_names.append(source["name"])
            if serving.from_product_name:
                parsed_serving_count += 1
            for nutrient_key, (_, unit) in NUTRIENT_COLUMNS.items():
                value, quality, nutrient_source, source_ref = None, "MISSING", "NONE", None
                if mfds and mfds.nutrients_per_basis.get(nutrient_key) is not None:
                    value = mfds.nutrients_per_basis[nutrient_key] * serving.amount / mfds.basis_amount
                    quality, nutrient_source, source_ref = "CONFIRMED", "MFDS", mfds.item_report_no
                elif standard and standard["nutrients"].get(nutrient_key) is not None:
                    value = standard["nutrients"][nutrient_key] * serving.amount / 100
                    quality, nutrient_source, source_ref = "ESTIMATED", "NATIONAL_STANDARD", standard_code
                db.add(FoodNutrient(food_id=food.id, nutrient_key=nutrient_key,
                                    value=Decimal(str(round(value, 4))) if value is not None else None,
                                    unit=unit, quality=quality, source=nutrient_source, source_ref=source_ref))

        db.commit()
        candidate_count = sum(bool(row["item_report_candidates"]) for row in cu_rows.values())
        print(f"CU 상품 {len(cu_rows)}개를 적재했습니다.")
        print(f"품목제조보고번호 후보가 있는 상품: {candidate_count}개")
        print(f"국가표준 또는 식약처 영양 연결 상품: {len(enriched_names)}개")
        print(f"상품명에서 판매 용량을 추출한 상품: {parsed_serving_count}개")


def main() -> None:
    parser = argparse.ArgumentParser(description="CU 테스트 상품과 영양 데이터를 적재합니다.")
    parser.add_argument("--cu", required=True, type=existing_file)
    parser.add_argument("--standard", required=True, type=existing_file)
    args = parser.parse_args()
    seed(args.cu, args.standard)


if __name__ == "__main__":
    main()
